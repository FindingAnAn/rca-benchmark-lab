"""Extract schema only from locally supplied documents, never sample subscriber values."""
import argparse
import hashlib
import json
from pathlib import Path


def run(doc,book,output):
    from docx import Document
    from openpyxl import load_workbook
    d=Document(doc);w=load_workbook(book,read_only=True,data_only=True)
    aliases={'CDR_Recurring_success':'Recurring_success_push','CDR_Recurring_fail':'Recurring_fail_push',
             'IPTV_recurring_by_OCS':'IPTV recurring by OCS'}
    items=[]
    for r in d.tables[0].rows:
        c=[x.text.strip() for x in r.cells]
        if not c[0].isdigit():continue
        name=c[2];matches=[s for s in w.sheetnames if s.lower()==name.lower()]
        if not matches:matches=[s for s in w.sheetnames if s.lower().startswith(name.lower())]
        if not matches:matches=[aliases[name]]
        s=w[matches[0]];rows=list(s.values);idx=list(rows[0]).index('prop_name')
        fields=[str(r[idx]) for r in rows[1:] if r[idx] is not None]
        items.append(dict(cdr_type=name,service_id=c[1],cycle_document=c[5],sheet=s.title,fields=fields,
            unique_fields=len(set(fields)),duplicate_field_names=sorted({f for f in fields if fields.count(f)>1}),
            schema_status='reference_requires_runtime_version_confirmation'))
    out=Path(output);out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(dict(types=items,sources={Path(x).name:hashlib.sha256(Path(x).read_bytes()).hexdigest() for x in [doc,book]},
        note='Schema only; no raw subscriber examples. Duplicate prop_name rows are not resolved by guessing.'),ensure_ascii=False,indent=2),encoding='utf-8')
    return [(r['cdr_type'],r['unique_fields']) for r in items]


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--doc',required=True);p.add_argument('--xlsx',required=True);p.add_argument('--output',required=True)
    a=p.parse_args();print(run(a.doc,a.xlsx,a.output))
